#!/usr/bin/env python3
"""Renders rapport-securite.md (+ analyse-critique.md as an annex) into an
editable Word document using python-docx. Small purpose-built markdown
converter (headings, tables, bold/code inline spans, lists, fenced code,
italics) mirroring build_pdf.py's parser, targeting python-docx instead of
reportlab so the output uses real Word heading styles / tables / lists."""
import re
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
MAIN_MD = os.path.join(HERE, "rapport-securite.md")
ANNEX_MD = os.path.join(HERE, "analyse-critique.md")
OUT_DOCX = os.path.join(HERE, "rapport-securite.docx")

HEADER_BLUE = RGBColor(0x2B, 0x3A, 0x67)
BAND_GREY = "F2F2F7"

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline_runs(paragraph, text, base_size=None):
    """Split text on **bold** / `code` spans and add correctly formatted runs."""
    for token in TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(token)
        if base_size:
            run.font.size = base_size


def shade_cell(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_page_numbers(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    run = p.add_run("Examen final — Sécurité des données — Ibrahima Lo — Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run2 = p.add_run()
    run2.font.size = Pt(8)
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_sep)
    run2._r.append(fld_end)


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(row)
        i += 1
    if len(rows) >= 2:
        del rows[1]
    return rows, i


def add_table(doc, rows):
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            para = cell.paragraphs[0]
            add_inline_runs(para, cell_text, base_size=Pt(9))
            if r == 0:
                shade_cell(cell, "2B3A67")
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True
            elif r % 2 == 0:
                shade_cell(cell, BAND_GREY)
    doc.add_paragraph()


def build(doc, lines, skip_h1=False):
    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("# "):
            if not skip_h1:
                h = doc.add_heading(stripped[2:], level=0)
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped)
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if re.match(r"^\d+\.\s", stripped):
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                txt = re.sub(r"^\d+\.\s", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                add_inline_runs(p, txt)
                i += 1
            continue

        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                txt = lines[i].strip()[2:]
                p = doc.add_paragraph(style="List Bullet")
                add_inline_runs(p, txt)
                i += 1
            continue

        if stripped.startswith("_") and stripped.endswith("_") and len(stripped) > 2:
            p = doc.add_paragraph()
            run = p.add_run(stripped[1:-1])
            run.italic = True
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**") and len(stripped) > 2:
            p = doc.add_paragraph()
            run = p.add_run(stripped[1:-1])
            run.italic = True
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, stripped)
        i += 1


def read_lines(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read().splitlines()


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    main_lines = read_lines(MAIN_MD)
    annex_lines = read_lines(ANNEX_MD)

    split_at = next(idx for idx, ln in enumerate(main_lines) if ln.strip().startswith("## 1."))
    header_lines, body_lines = main_lines[:split_at], main_lines[split_at:]

    build(doc, header_lines)

    doc.add_page_break()
    doc.add_heading("Sommaire", level=1)
    for ln in main_lines:
        if ln.strip().startswith("## "):
            doc.add_paragraph(ln.strip()[3:], style="List Bullet")
    doc.add_paragraph("Annexe — Partie 6 : Analyse critique", style="List Bullet")
    doc.add_page_break()

    build(doc, body_lines)

    doc.add_page_break()
    doc.add_heading("Annexe — Partie 6 : Analyse critique", level=0)
    build(doc, annex_lines, skip_h1=True)

    add_page_numbers(doc)

    doc.save(OUT_DOCX)
    print(f"Written: {OUT_DOCX}")


if __name__ == "__main__":
    main()
