#!/usr/bin/env python3
"""Builds the oral-defense slide deck (presentation-soutenance.pptx) with
python-pptx (no Node/pptxgenjs available on this machine). Self-contained
content — not a markdown->slides converter — mirroring the report's real
numbers and structure."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "presentation-soutenance.pptx")

# ---- palette: "Midnight Executive" + severity accents -----------------
NAVY = RGBColor(0x1E, 0x27, 0x61)
NAVY_DARK = RGBColor(0x12, 0x18, 0x3D)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x26, 0x3A)
MUTED = RGBColor(0x6B, 0x72, 0x8E)
CARD_BG = RGBColor(0xF3, 0xF5, 0xFB)
CRITICAL = RGBColor(0xC0, 0x1E, 0x2B)
HIGH = RGBColor(0xE0, 0x7A, 0x1F)
MEDIUM = RGBColor(0xD8, 0xA3, 0x00)
GOOD = RGBColor(0x1E, 0x8A, 0x5F)

FONT_HEAD = "Cambria"
FONT_BODY = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = color
    rect.line.fill.background()
    rect.shadow.inherit = False
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return rect


def textbox(slide, l, t, w, h, text, size=16, color=INK, bold=False, italic=False,
            align=PP_ALIGN.LEFT, font=FONT_BODY, anchor=None, line_spacing=None):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def rect(slide, l, t, w, h, color, radius=False, line=False, line_color=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, l, t, w, h)
    if radius:
        try:
            sh.adjustments[0] = 0.06
        except Exception:
            pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line:
        sh.line.color.rgb = line_color or color
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def circle(slide, l, t, d, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, d, d)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def page_footer(slide, page_no, dark=False):
    color = RGBColor(0xB8, 0xC2, 0xE8) if dark else MUTED
    textbox(slide, Inches(0.5), Inches(7.08), Inches(9), Inches(0.35),
            "Examen final — Sécurité des données — Ibrahima Lo", size=10, color=color)
    textbox(slide, Inches(12.3), Inches(7.08), Inches(0.6), Inches(0.35),
            str(page_no), size=10, color=color, align=PP_ALIGN.RIGHT)


def severity_chip(slide, l, t, label, color):
    w, h = Inches(1.1), Inches(0.32)
    chip = rect(slide, l, t, w, h, color, radius=True)
    tf = chip.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = FONT_BODY
    return chip


# =========================================================================
# Slide 1 — Title
# =========================================================================
def slide_title(prs):
    s = blank(prs)
    bg(s, NAVY_DARK)
    textbox(s, Inches(1), Inches(2.3), Inches(11.3), Inches(0.6),
            "AUDIT DE SÉCURITÉ APPLICATIVE", size=18, color=ICE, bold=True, font=FONT_BODY)
    textbox(s, Inches(1), Inches(2.85), Inches(11.3), Inches(1.8),
            "OWASP Juice Shop : identification, remédiation\net automatisation DevSecOps",
            size=38, color=WHITE, bold=True, font=FONT_HEAD, line_spacing=1.05)
    textbox(s, Inches(1), Inches(4.9), Inches(10), Inches(0.5),
            "Examen final — Sécurité des données · Licence 3 Cybersécurité", size=16, color=ICE)
    textbox(s, Inches(1), Inches(6.2), Inches(8), Inches(0.5),
            "Ibrahima Lo — Cybersecurity Analyst / Junior DevSecOps Engineer",
            size=14, color=WHITE, bold=True)


# =========================================================================
# Slide 2 — Agenda
# =========================================================================
def slide_agenda(prs):
    s = blank(prs)
    bg(s, WHITE)
    rect(s, 0, 0, Inches(4.2), SH, NAVY)
    textbox(s, Inches(0.5), Inches(0.7), Inches(3.3), Inches(1),
            "SOMMAIRE", size=26, color=WHITE, bold=True, font=FONT_HEAD)
    textbox(s, Inches(0.5), Inches(1.6), Inches(3.4), Inches(4.5),
            "Une démarche d'audit\ncomplète : identifier,\nprioriser, corriger,\nautomatiser, décider.",
            size=15, color=ICE, italic=True, line_spacing=1.3)

    items = [
        "Contexte et scénario",
        "Méthodologie",
        "Vulnérabilités identifiées (V1–V6)",
        "Analyse d'impact CWE / CIA",
        "Remédiation et vérification",
        "Pipeline Jenkins DevSecOps",
        "Résultats avant / après",
        "Décision de déploiement",
        "Analyse critique",
    ]
    top = Inches(0.7)
    for i, item in enumerate(items):
        y = top + Inches(0.62) * i
        circle(s, Inches(4.7), y, Inches(0.42), NAVY if i % 2 == 0 else ICE)
        num_color = WHITE if i % 2 == 0 else NAVY
        tb = textbox(s, Inches(4.7), y, Inches(0.42), Inches(0.42), str(i + 1),
                     size=15, color=num_color, bold=True, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, Inches(5.3), y + Inches(0.02), Inches(7.3), Inches(0.42),
                item, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    page_footer(s, 2)


# =========================================================================
# Slide 3 — Contexte
# =========================================================================
def slide_context(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.5), Inches(10), Inches(0.7),
            "Contexte et scénario", size=32, color=NAVY, bold=True, font=FONT_HEAD)
    rect(s, Inches(0.6), Inches(1.35), Inches(7.6), Inches(4.6), CARD_BG, radius=True)
    textbox(s, Inches(0.95), Inches(1.65), Inches(6.9), Inches(0.5),
            "La mission", size=16, color=NAVY, bold=True)
    textbox(s, Inches(0.95), Inches(2.15), Inches(6.9), Inches(3.7),
            "Une organisation prépare la mise en production d'une application web "
            "permettant la création de compte, l'authentification et l'accès à des "
            "ressources contenant des données sensibles (produits, paniers, "
            "commandes, profils).\n\n"
            "Rôle confié : Cybersecurity Analyst / Junior DevSecOps Engineer — "
            "évaluer la sécurité, analyser les risques, corriger, puis intégrer des "
            "contrôles automatisés avant toute décision de déploiement.",
            size=15, color=INK, line_spacing=1.25)

    rect(s, Inches(8.5), Inches(1.35), Inches(4.25), Inches(4.6), NAVY, radius=True)
    textbox(s, Inches(8.8), Inches(1.6), Inches(3.7), Inches(0.5),
            "Application cible", size=15, color=ICE, bold=True)
    textbox(s, Inches(8.8), Inches(2.1), Inches(3.7), Inches(0.6),
            "OWASP Juice Shop", size=20, color=WHITE, bold=True, font=FONT_HEAD)
    textbox(s, Inches(8.8), Inches(2.75), Inches(3.7), Inches(2.9),
            "• Clone local, patché\n"
            "• Déployée via Docker Desktop\n"
            "• Auditée en conditions réelles\n"
            "  (machine physique, 8 Go RAM)\n"
            "• Pipeline CI/CD Jenkins",
            14, color=WHITE, line_spacing=1.4)
    page_footer(s, 3)


# =========================================================================
# Slide 4 — Méthodologie
# =========================================================================
def slide_methodology(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.5), Inches(10), Inches(0.7),
            "Méthodologie", size=32, color=NAVY, bold=True, font=FONT_HEAD)

    cards = [
        ("Revue manuelle", "Lecture ciblée du code : authentification, panier, module de sécurité, frontend Angular"),
        ("SAST — Semgrep", "Règles OWASP Top 10 / JS / TS sur l'ensemble du code source"),
        ("SCA — Trivy", "Analyse des dépendances npm (package-lock.json) pour les CVE connues"),
        ("Secrets — Gitleaks", "Recherche de clés, tokens et mots de passe codés en dur"),
        ("DAST — OWASP ZAP", "Scan dynamique de l'application en fonctionnement (baseline)"),
        ("Analyse CIA", "Confidentialité / Intégrité / Disponibilité pour chaque vulnérabilité"),
    ]
    cols, rows = 3, 2
    cw, ch = Inches(3.95), Inches(2.15)
    gx, gy = Inches(0.25), Inches(0.25)
    x0, y0 = Inches(0.6), Inches(1.5)
    for i, (title, desc) in enumerate(cards):
        r, c = divmod(i, cols)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        rect(s, x, y, cw, ch, CARD_BG, radius=True)
        circle(s, x + Inches(0.25), y + Inches(0.25), Inches(0.5), NAVY)
        textbox(s, x + Inches(0.25), y + Inches(0.25), Inches(0.5), Inches(0.5),
                str(i + 1), size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x + Inches(0.9), y + Inches(0.22), cw - Inches(1.1), Inches(0.55),
                title, size=15, color=NAVY, bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(0.25), y + Inches(0.85), cw - Inches(0.5), Inches(1.2),
                desc, size=12, color=INK, line_spacing=1.2)
    page_footer(s, 4)


# =========================================================================
# Slide 5 — Vulnérabilités (table)
# =========================================================================
def slide_vulns(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "6 vulnérabilités identifiées", size=30, color=NAVY, bold=True, font=FONT_HEAD)

    rows = [
        ("V1", "SQL Injection", "routes/login.ts", "CWE-89", "Critical", CRITICAL),
        ("V2", "XSS (DOM-based)", "search-result.component.ts", "CWE-79", "High", HIGH),
        ("V3", "Broken Access Control (IDOR)", "routes/basket.ts", "CWE-639", "Critical", CRITICAL),
        ("V4", "Secret codé en dur (clé JWT)", "lib/insecurity.ts", "CWE-798", "Critical", CRITICAL),
        ("V5", "Hash mot de passe faible (MD5)", "lib/insecurity.ts", "CWE-916", "Medium", MEDIUM),
        ("V6", "Dépendances vulnérables", "package.json", "CWE-1104", "High", HIGH),
    ]
    x0, y0 = Inches(0.6), Inches(1.4)
    col_w = [Inches(0.7), Inches(3.5), Inches(3.5), Inches(1.6), Inches(1.9)]
    row_h = Inches(0.78)
    headers = ["ID", "Vulnérabilité", "Composant", "CWE", "Criticité"]
    for c, htext in enumerate(headers):
        x = x0 + sum(col_w[:c])
        rect(s, x, y0, col_w[c], Inches(0.5), NAVY)
        textbox(s, x + Inches(0.1), y0, col_w[c] - Inches(0.2), Inches(0.5), htext,
                size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    for r, (vid, name, comp, cwe, sev, color) in enumerate(rows):
        y = y0 + Inches(0.5) + r * row_h
        band = CARD_BG if r % 2 == 0 else WHITE
        rect(s, x0, y, sum(col_w, Emu(0)), row_h, band)
        vals = [vid, name, comp, cwe]
        for c, val in enumerate(vals):
            x = x0 + sum(col_w[:c])
            bold = c == 0
            textbox(s, x + Inches(0.1), y, col_w[c] - Inches(0.2), row_h, val,
                    size=13, color=INK, bold=bold, anchor=MSO_ANCHOR.MIDDLE)
        chip_x = x0 + sum(col_w[:4]) + Inches(0.35)
        severity_chip(s, chip_x, y + Inches(0.22), sev, color)
    page_footer(s, 5)


# =========================================================================
# Slide 6 — Impact CIA
# =========================================================================
def slide_cia(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "Impact sur la sécurité des données (CIA)", 28,
            color=NAVY, bold=True, font=FONT_HEAD)

    cats = [
        ("Confidentialité", "5 / 6 vulnérabilités permettent un accès non autorisé aux données", ICE),
        ("Intégrité", "3 / 6 permettent une modification non autorisée des données", ICE),
        ("Disponibilité", "Impact indirect (requêtes malformées, comptes admin forgés)", ICE),
    ]
    x0 = Inches(0.6)
    cw = Inches(4.0)
    for i, (label, desc, color) in enumerate(cats):
        x = x0 + i * (cw + Inches(0.2))
        rect(s, x, Inches(1.4), cw, Inches(2.6), NAVY, radius=True)
        textbox(s, x + Inches(0.3), Inches(1.65), cw - Inches(0.6), Inches(0.6),
                label, size=18, color=WHITE, bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(0.3), Inches(2.35), cw - Inches(0.6), Inches(1.5),
                desc, size=13, color=ICE, line_spacing=1.3)

    textbox(s, Inches(0.6), Inches(4.35), Inches(12), Inches(0.5),
            "Criticité et priorité de correction", size=18, color=NAVY, bold=True, font=FONT_HEAD)
    stats = [("3", "vulnérabilités\nCritical", CRITICAL), ("2", "vulnérabilités\nHigh", HIGH),
             ("1", "vulnérabilité\nMedium", MEDIUM)]
    for i, (num, label, color) in enumerate(stats):
        x = Inches(0.6) + i * Inches(2.3)
        textbox(s, x, Inches(4.9), Inches(1.8), Inches(1.1), num, size=54, color=color,
                bold=True, font=FONT_HEAD)
        textbox(s, x, Inches(5.95), Inches(2.2), Inches(0.7), label, size=13, color=INK,
                line_spacing=1.1)
    textbox(s, Inches(7.5), Inches(4.9), Inches(5.2), Inches(1.9),
            "V1, V3 et V4 permettent chacune, seules, une compromission "
            "totale des données sans privilège préalable → correction "
            "immédiate obligatoire avant toute mise en production.",
            size=14, color=INK, italic=True, line_spacing=1.3)
    page_footer(s, 6)


# =========================================================================
# Slide 7 — Remédiation approche
# =========================================================================
def slide_remediation(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "Remédiation : 3 correctifs critiques", size=30, color=NAVY, bold=True, font=FONT_HEAD)

    items = [
        ("V1 — SQL Injection", "Cause", "Concaténation de chaînes dans la requête SQL",
         "Correction", "Requête paramétrée (Sequelize `replacements`)"),
        ("V3 — IDOR (panier)", "Cause", "Aucune vérification que le panier appartient à l'utilisateur",
         "Correction", "Contrôle serveur : id demandé == bid du token, sinon 403"),
        ("V4 — Secret JWT codé en dur", "Cause", "Clé privée RSA en clair, versionnée dans le dépôt",
         "Correction", "Rotation de clé + chargement depuis un fichier exclu de Git"),
    ]
    y0 = Inches(1.4)
    for i, (title, l1, v1, l2, v2) in enumerate(items):
        y = y0 + i * Inches(1.75)
        rect(s, Inches(0.6), y, Inches(12.1), Inches(1.55), CARD_BG, radius=True)
        circle(s, Inches(0.95), y + Inches(0.58), Inches(0.4), GOOD)
        textbox(s, Inches(0.95), y + Inches(0.58), Inches(0.4), Inches(0.4), "✓",
                size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, Inches(1.55), y + Inches(0.12), Inches(3.0), Inches(1.3),
                title, size=16, color=NAVY, bold=True, font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, Inches(4.7), y + Inches(0.15), Inches(0.9), Inches(0.35), l1,
                size=11, color=MUTED, bold=True)
        textbox(s, Inches(4.7), y + Inches(0.5), Inches(3.5), Inches(0.9), v1,
                size=12.5, color=INK, line_spacing=1.15)
        textbox(s, Inches(8.6), Inches(0) + y + Inches(0.15), Inches(1.2), Inches(0.35), l2,
                size=11, color=MUTED, bold=True)
        textbox(s, Inches(8.6), y + Inches(0.5), Inches(3.9), Inches(0.9), v2,
                size=12.5, color=INK, line_spacing=1.15)
    page_footer(s, 7)


# =========================================================================
# Slide 8 — Preuves avant / après
# =========================================================================
def slide_before_after(prs):
    s = blank(prs)
    bg(s, NAVY_DARK)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "Vérification : avant / après correction", size=28, color=WHITE, bold=True,
            font=FONT_HEAD)

    data = [
        ("Semgrep", "9", "7", "findings"),
        ("Gitleaks", "67", "66", "secrets détectés"),
    ]
    x0 = Inches(0.6)
    cw = Inches(5.9)
    for i, (tool, before, after, unit) in enumerate(data):
        x = x0 + i * (cw + Inches(0.3))
        rect(s, x, Inches(1.4), cw, Inches(2.3), RGBColor(0x24, 0x2E, 0x70), radius=True)
        textbox(s, x + Inches(0.35), Inches(1.6), cw - Inches(0.7), Inches(0.5),
                tool, size=17, color=ICE, bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(0.35), Inches(2.15), Inches(2), Inches(1.1), before,
                size=48, color=RGBColor(0xE8, 0x8A, 0x8A), bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(2.5), Inches(2.55), Inches(0.6), Inches(0.5), "→",
                size=28, color=WHITE)
        textbox(s, x + Inches(3.1), Inches(2.15), Inches(2), Inches(1.1), after,
                size=48, color=RGBColor(0x8C, 0xE0, 0xB0), bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(0.35), Inches(3.2), cw - Inches(0.7), Inches(0.4),
                unit, size=12, color=ICE)

    rect(s, Inches(0.6), Inches(4.05), Inches(12.1), Inches(2.6), RGBColor(0x24, 0x2E, 0x70), radius=True)
    textbox(s, Inches(0.95), Inches(4.25), Inches(11.4), Inches(0.5),
            "Tests manuels en direct (logique métier — non couverte par les outils)",
            size=16, color=ICE, bold=True, font=FONT_HEAD)
    textbox(s, Inches(0.95), Inches(4.85), Inches(5.4), Inches(1.6),
            "SQLi (V1)\n"
            "Payload ' OR 1=1-- réussissait avant\ncorrection → échoue désormais.",
            size=13.5, color=WHITE, line_spacing=1.35)
    textbox(s, Inches(6.7), Inches(4.85), Inches(5.6), Inches(1.6),
            "IDOR (V3)\n"
            "fetch('/rest/basket/2') renvoyait le panier\nd'un autre utilisateur → 403 Forbidden désormais.",
            size=13.5, color=WHITE, line_spacing=1.35)
    page_footer(s, 8, dark=True)


# =========================================================================
# Slide 9 — Pipeline Jenkins
# =========================================================================
def slide_pipeline(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "Pipeline Jenkins DevSecOps", 30, color=NAVY, bold=True,
            font=FONT_HEAD)

    steps = ["Checkout", "Build /\nPreparation", "Security\nAnalysis", "DAST\n(ZAP)",
             "Report\nGeneration", "Notification"]
    n = len(steps)
    total_w = Inches(12.1)
    gap = Inches(0.2)
    box_w = Emu(int((total_w - gap * (n - 1)) / n))
    x0 = Inches(0.6)
    y = Inches(1.6)
    box_h = Inches(1.3)
    for i, step in enumerate(steps):
        x = x0 + i * (box_w + gap)
        color = NAVY if i % 2 == 0 else RGBColor(0x2E, 0x3B, 0x8C)
        rect(s, x, y, box_w, box_h, color, radius=True)
        textbox(s, x + Inches(0.08), y, box_w - Inches(0.16), box_h, step, size=13,
                color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                line_spacing=1.1)
        if i < n - 1:
            textbox(s, x + box_w, y, gap, box_h, "→", size=18, color=NAVY,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    textbox(s, Inches(0.6), Inches(3.3), Inches(12), Inches(0.5),
            "3 scanners au stage Security Analysis (séquentiel — contrainte mémoire de la machine)",
            size=14, color=MUTED, italic=True)

    tools = [
        ("Semgrep", "SAST", "Injections, API dangereuses (innerHTML, eval)"),
        ("Trivy", "SCA", "CVE connues dans les dépendances npm"),
        ("Gitleaks", "Secrets", "Clés API, tokens, mots de passe codés en dur"),
        ("OWASP ZAP", "DAST", "En-têtes manquants, config HTTP, XSS reflété"),
    ]
    cw = Inches(2.95)
    for i, (name, kind, desc) in enumerate(tools):
        x = Inches(0.6) + i * (cw + Inches(0.15))
        rect(s, x, Inches(3.95), cw, Inches(2.6), CARD_BG, radius=True)
        severity_chip(s, x + Inches(0.2), Inches(4.15), kind, NAVY)
        textbox(s, x + Inches(0.2), Inches(4.6), cw - Inches(0.4), Inches(0.4),
                name, size=15, color=NAVY, bold=True, font=FONT_HEAD)
        textbox(s, x + Inches(0.2), Inches(5.05), cw - Inches(0.4), Inches(1.4),
                desc, size=12, color=INK, line_spacing=1.2)
    page_footer(s, 9)


# =========================================================================
# Slide 10 — Résultats
# =========================================================================
def slide_results(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.7),
            "Résultats du pipeline (après remédiation)", size=28, color=NAVY, bold=True,
            font=FONT_HEAD)

    rows = [
        ("Semgrep (SAST)", "7 findings (2 error / 5 warning)"),
        ("Trivy (SCA)", "8 critical / 45 high / 33 medium / 6 low"),
        ("Gitleaks (secrets)", "66 secrets détectés (essentiellement des faux positifs de test)"),
        ("OWASP ZAP (DAST)", "Medium: 2, Low: 5, Informational: 3"),
    ]
    x0, y0 = Inches(0.6), Inches(1.5)
    row_h = Inches(0.85)
    for i, (tool, val) in enumerate(rows):
        y = y0 + i * row_h
        band = CARD_BG if i % 2 == 0 else WHITE
        rect(s, x0, y, Inches(12.1), row_h, band)
        textbox(s, x0 + Inches(0.3), y, Inches(3.5), row_h, tool, size=15, color=NAVY,
                bold=True, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x0 + Inches(4), y, Inches(8), row_h, val, size=14, color=INK,
                anchor=MSO_ANCHOR.MIDDLE)

    textbox(s, Inches(0.6), Inches(5.1), Inches(12), Inches(0.5),
            "Classement par priorité", size=18, color=NAVY, bold=True, font=FONT_HEAD)
    chips = [("SQLi — Bloquer", CRITICAL), ("IDOR — Bloquer", CRITICAL),
             ("Secret JWT — Bloquer", CRITICAL), ("XSS — Requis", HIGH),
             ("Dépendances — Requis", HIGH), ("Hash MD5 — Planifié", MEDIUM)]
    x = Inches(0.6)
    for label, color in chips:
        w = Inches(0.32 * len(label) / 10 + 1.4)
        chip = rect(s, x, Inches(5.65), w, Inches(0.5), color, radius=True)
        tf = chip.text_frame
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = WHITE
        x += w + Inches(0.15)
    page_footer(s, 10)


# =========================================================================
# Slide 11 — Décision de déploiement
# =========================================================================
def slide_decision(prs):
    s = blank(prs)
    bg(s, NAVY_DARK)
    textbox(s, Inches(0.6), Inches(0.5), Inches(11), Inches(0.7),
            "Décision de déploiement", size=30, color=WHITE, bold=True, font=FONT_HEAD)

    rect(s, Inches(0.8), Inches(1.7), Inches(5.6), Inches(2.4), CRITICAL, radius=True)
    textbox(s, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5), "AVANT remédiation",
            size=15, color=WHITE, bold=True)
    textbox(s, Inches(1.1), Inches(2.5), Inches(5), Inches(1.0), "Reject\nDeployment",
            size=30, color=WHITE, bold=True, font=FONT_HEAD, line_spacing=1.0)

    rect(s, Inches(6.9), Inches(1.7), Inches(5.6), Inches(2.4), GOOD, radius=True)
    textbox(s, Inches(7.2), Inches(1.95), Inches(5), Inches(0.5), "APRÈS remédiation",
            size=15, color=WHITE, bold=True)
    textbox(s, Inches(7.2), Inches(2.5), Inches(5), Inches(1.0), "Accept with\nConditions",
            size=30, color=WHITE, bold=True, font=FONT_HEAD, line_spacing=1.0)

    textbox(s, Inches(0.8), Inches(4.4), Inches(11.7), Inches(2.4),
            "Trois vulnérabilités critiques (V1, V3, V4) permettaient une compromission "
            "totale de la confidentialité et de l'intégrité des données sans privilège "
            "préalable — un déploiement en l'état aurait exposé l'organisation à une fuite "
            "massive et à une usurpation de compte administrateur.\n\n"
            "Après correction et vérification, ces risques sont neutralisés. Les "
            "vulnérabilités restantes (V2, V5, V6) sont acceptables temporairement, à "
            "condition d'être planifiées et corrigées, avec un contrôle de non-régression "
            "assuré par le pipeline Jenkins.",
            size=15, color=ICE, line_spacing=1.35)
    page_footer(s, 11, dark=True)


# =========================================================================
# Slide 12 — Analyse critique
# =========================================================================
def slide_critical(prs):
    s = blank(prs)
    bg(s, WHITE)
    textbox(s, Inches(0.6), Inches(0.5), Inches(11), Inches(0.7),
            "Analyse critique", size=30, color=NAVY, bold=True, font=FONT_HEAD)
    textbox(s, Inches(0.6), Inches(1.3), Inches(11.5), Inches(0.6),
            "Un outil de sécurité qui ne détecte aucune vulnérabilité "
            "peut-il garantir qu'une application est sécurisée ?",
            size=18, color=MUTED, italic=True, line_spacing=1.2)

    points = [
        ("Faux négatifs", "Chaque outil ne détecte que ce que ses règles couvrent — l'IDOR (V3) n'aurait "
                           "jamais été trouvée par Semgrep seul : c'est un défaut de logique métier."),
        ("Faux positifs", "Une configuration trop permissive ou mal calée masque de vrais problèmes et "
                           "génère de l'« alert fatigue » (ex. les ~66 alertes Gitleaks sur des données de test)."),
        ("Couverture limitée", "Un SAST/DAST/SCA couvre des catégories connues (OWASP Top 10, CWE) — la "
                                "surface de risque réelle d'une application les dépasse toujours."),
        ("Analyse humaine", "Le pipeline automatisé garantit la reproductibilité et empêche la régression, "
                             "mais doit être complété par une revue de code et une analyse de risques ciblée."),
    ]
    y0 = Inches(2.3)
    for i, (title, desc) in enumerate(points):
        r, c = divmod(i, 2)
        x = Inches(0.6) + c * Inches(6.15)
        y = y0 + r * Inches(2.25)
        rect(s, x, y, Inches(5.9), Inches(2.05), CARD_BG, radius=True)
        circle(s, x + Inches(0.3), y + Inches(0.28), Inches(0.5), NAVY)
        textbox(s, x + Inches(0.3), y + Inches(0.28), Inches(0.5), Inches(0.5), str(i + 1),
                size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x + Inches(1.0), y + Inches(0.22), Inches(4.7), Inches(0.6),
                title, size=16, color=NAVY, bold=True, font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x + Inches(0.3), y + Inches(0.95), Inches(5.35), Inches(1.0),
                desc, size=12.5, color=INK, line_spacing=1.25)
    page_footer(s, 12)


# =========================================================================
# Slide 13 — Conclusion
# =========================================================================
def slide_conclusion(prs):
    s = blank(prs)
    bg(s, NAVY)
    textbox(s, Inches(0.8), Inches(0.7), Inches(11), Inches(0.7),
            "Conclusion", size=32, color=WHITE, bold=True, font=FONT_HEAD)

    items = [
        "6 vulnérabilités identifiées, couvrant l'OWASP Top 10 (injection, XSS, "
        "contrôle d'accès, secrets, cryptographie faible, dépendances vulnérables)",
        "3 corrections critiques appliquées et vérifiées, par outil ET par test manuel",
        "Pipeline Jenkins opérationnel : SAST + SCA + Secret Detection + DAST",
        "Un audit ponctuel devient un garde-fou continu contre la régression",
        "Les outils automatisés réduisent le risque — ils ne remplacent pas le jugement humain",
    ]
    y = Inches(2.0)
    for item in items:
        circle(s, Inches(0.9), y + Inches(0.05), Inches(0.28), ICE)
        textbox(s, Inches(1.5), y, Inches(10.8), Inches(0.85), item, size=16, color=WHITE,
                line_spacing=1.2, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.95)
    page_footer(s, 13, dark=True)


# =========================================================================
# Slide 14 — Merci / lien
# =========================================================================
def slide_thanks(prs):
    s = blank(prs)
    bg(s, NAVY_DARK)
    textbox(s, Inches(1), Inches(2.3), Inches(11), Inches(1), "Merci de votre attention",
            size=40, color=WHITE, bold=True, font=FONT_HEAD)
    textbox(s, Inches(1), Inches(3.3), Inches(10), Inches(0.5), "Questions ?", size=20,
            color=ICE, italic=True)

    rect(s, Inches(1), Inches(4.6), Inches(11.3), Inches(1.3), RGBColor(0x24, 0x2E, 0x70),
         radius=True)
    textbox(s, Inches(1.35), Inches(4.8), Inches(10.6), Inches(0.4),
            "Dépôt Git complet (code, pipeline, rapports, preuves) :", size=13, color=ICE)
    textbox(s, Inches(1.35), Inches(5.2), Inches(10.6), Inches(0.55),
            "https://github.com/ibrahima143lo-arch/examen-securite-donnees-ibrahima-lo",
            size=18, color=WHITE, bold=True)
    page_footer(s, 14, dark=True)


def main():
    prs = new_deck()
    slide_title(prs)
    slide_agenda(prs)
    slide_context(prs)
    slide_methodology(prs)
    slide_vulns(prs)
    slide_cia(prs)
    slide_remediation(prs)
    slide_before_after(prs)
    slide_pipeline(prs)
    slide_results(prs)
    slide_decision(prs)
    slide_critical(prs)
    slide_conclusion(prs)
    slide_thanks(prs)
    prs.save(OUT)
    print(f"Written: {OUT} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
